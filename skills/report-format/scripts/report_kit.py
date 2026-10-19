#!/usr/bin/env python3
r"""report_kit: 주제와 무관한 기술 보고서 Word(.docx) 빌더 (report.pdf 형식을 Word 로 옮긴 것).

Word 는 python-docx 로 만든다. PDF 가 필요하면 Word 에서 "PDF 로 저장"한다(이 빌더는 PDF 를 만들지 않는다).
미리보기는 설치된 Microsoft Word(Windows·WSL) 또는 LibreOffice 로 쪽 PNG 를 그린다. 둘 다 없으면 typst 근사 미리보기.
필요한 것: Python 3 + `pip install python-docx typst`. 그림을 그리는 예제 스크립트만 matplotlib 을 쓴다(선택).

  1) markdown 파일:   python3 report_kit.py build report.md [-o out.docx] [--preview [DIR]]
  2) Python API:      from report_kit import Report, num; r = Report(title, subtitle); r.h2(...); r.table(...); r.build("out.docx")
  미리보기:           python3 report_kit.py preview report.docx|report.md [--dir DIR] [--ppi 55]   → 쪽별 PNG + 접촉 시트(4쪽씩)
  환경 점검:          python3 report_kit.py doctor [--build]                          → python-docx·typst·폰트·Word 상태(--build: 예제를 실제로 빌드)
  환경 준비:          python3 report_kit.py setup [--no-pip] [--no-font]              → python-docx·typst 설치(pip → 사용자 폴더 → 예비 venv 순) + 한글 폰트(~/.local/share/report-format/fonts)

이 스크립트는 스킬 폴더(report-format/) 안 상대 위치만 쓴다. 폴더째 어디로 옮겨도 그대로 동작한다.

markdown 문법(최소 집합, 외부 패키지 불필요):
  ---                      머리말: title / subtitle / output / font (선택)
  title: 제목
  subtitle: 대상 · 범위 · 2026-09-15
  ---
  ## 1. 장 제목            → 장(왼쪽 세로선 + 회색 띠)     ### 1.1 절 제목 → 절(굵은 글자)
  :::key 핵심 결과          → 강조 상자(제목 뒤 불릿·문단), ::: 로 닫음. 문서 첫머리에 하나만
  :::two                   → 2단. 첫 단과 둘째 단은 ||| 한 줄로 나눔, ::: 로 닫음
  | a | b |  /  |---|---|  → 3선 표. 표 바로 다음 줄이 "표 N." 로 시작하면 표 아래 캡션(제목 한 줄)
  ![그림 1. 제목](path){width=66%}  → 그림 + 아래 캡션(제목 한 줄) (width 선택, 기본 100%)
  $$ 수식 $$                → 세리프 가운데 수식(한 줄). <sub>·<sup> 허용
  - 항목                    → 불릿
  \pagebreak               → 강제 쪽 나눔
  **굵게**, `코드`, <sub> <sup> <br> 허용. 문단은 빈 줄로 구분.

그림 자동 맞춤(빌드 시 기본 동작, references/rules.md W10):
  기준: 본문 폭 174 mm(2단 칸 약 84 mm)에서 좌우 빈 폭 ≤ 40 %, 높이 ≤ 쪽 본문의 75 %, 삽입 크기 ≥ 150 dpi.
  순서: 흰 여백 자르기(Pillow 있을 때) → 문제가 있으면 원천 재생성(regen) → 폭 조정(늘릴 때는 기준선까지만,
        이미지째 확대로 글자가 커지는 것을 막는다) → 세로로 긴 그림은 앞 문단 옆 배치 → 남은 것은 "미해결"로 출력.
  ![그림 1. 캡션](a.png){width=66% regen="python3 fig.py --w {w_mm} --h {h_mm} --out {out}"}   셸 명령 재생성
  ![그림 1. 캡션](a.png){fit=off}                                                          자동 맞춤 끄기
  Python API: r.fig(path, cap, regen=lambda w_mm, h_mm: make(w_mm, h_mm)) → 새 PNG 경로를 돌려주는 함수.
  미해결이 남으면 CLI 는 docx 를 쓴 뒤 종료 코드 3(--lenient 로 끔), API 는 build(strict=True) 또는 REPORT_STRICT=1 일 때 3.
  그래프는 figsize_mm(w_mm, h_mm) 로 실제 크기 그대로 그리면 글자 pt 가 지면 pt 와 같다.

숫자 표기(빌드 시 자동 점검, 경고):
  - 계산·측정값은 소수 셋째 자리까지 반올림(`1.234`), 천 단위 쉼표(`3,384 rpm`, `12,046 m³/hr`).
  - 정수 개수·절 번호·날짜(YYYY-MM-DD)·수식 안 상수·`코드` 는 예외. 점검이 잘못 잡으면 <span class="nocheck">…</span> 로 감싼다.
  - Python API 에서는 num(x) 로 형식을 만든다: num(3384) → '3,384', num(1.23456) → '1.235', num(0.5, 1) → '0.5'.
"""
import os, re, sys, glob, shutil, subprocess, argparse, tempfile, struct, zlib

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.normpath(os.path.join(HERE, ".."))
FONT_DIR = os.path.join(SKILL_DIR, "fonts")                       # 폴더째 복사한 사본에 폰트를 같이 둘 때
DATA_DIR = (os.path.join(os.environ.get("LOCALAPPDATA", os.path.expanduser("~")), "report-format") if os.name == "nt"
            else os.path.expanduser("~/.local/share/report-format"))   # setup 이 내려받는 것들의 자리. 플러그인이 갱신·재설치돼도 남는다
USER_FONT_DIR = os.path.join(DATA_DIR, "fonts")
VENV_DIR = os.path.join(DATA_DIR, "venv")     # 시스템 파이썬에 pip 설치가 막힌 PC(PEP 668)에서 setup 이 만드는 예비 환경
FONT_FAMILY = "Noto Sans KR"
FONT_FILES = {   # 정적(비가변) 파일. 가변 폰트(NotoSansKR[wght].ttf)는 typst 가 가장 가는 굵기로 읽으므로 쓰지 않는다.
    "NotoSansKR-Regular.otf": "https://github.com/notofonts/noto-cjk/raw/main/Sans/SubsetOTF/KR/NotoSansKR-Regular.otf",
    "NotoSansKR-Bold.otf": "https://github.com/notofonts/noto-cjk/raw/main/Sans/SubsetOTF/KR/NotoSansKR-Bold.otf",
}

# ---------------------------------------------------------------- 폰트
def _font_dirs(explicit=None):
    """폰트를 찾을 폴더 목록(우선순위 순). 인자·REPORT_FONT 는 파일이든 폴더든 받는다."""
    out = []
    for c in (explicit, os.environ.get("REPORT_FONT")):
        if c:
            c = os.path.expanduser(c); out.append(c if os.path.isdir(c) else os.path.dirname(c))
    out += [FONT_DIR, USER_FONT_DIR, os.path.expanduser("~/.fonts"), os.path.expanduser("~/.local/share/fonts"),
            os.path.expanduser("~/Library/Fonts"), "/usr/share/fonts", "/usr/local/share/fonts", r"C:\Windows\Fonts"]
    seen, res = set(), []
    for d in out:
        if d and os.path.isdir(d) and d not in seen:
            seen.add(d); res.append(d)
    return res

def _is_variable_font(path):
    try:
        with open(path, "rb") as f:
            return b"fvar" in f.read(4096)
    except OSError:
        return False

def font_status(explicit=None):
    """(정적 Regular 경로 또는 None, 정적 Bold 경로 또는 None, 가변 폰트 경로 또는 None, typst 에 넘길 폰트 폴더 목록)"""
    reg = bold = var = None
    for d in _font_dirs(explicit):
        for f in sorted(glob.glob(os.path.join(d, "**", "NotoSans*KR*"), recursive=True)):
            base = os.path.basename(f).lower()
            if not base.endswith((".otf", ".ttf")):
                continue
            if _is_variable_font(f):
                var = var or f; continue
            if "regular" in base and not reg:
                reg = f
            elif "bold" in base and not bold:
                bold = f
    use = sorted({os.path.dirname(reg), os.path.dirname(bold)}) if (reg and bold) else _font_dirs(explicit)
    return reg, bold, var, use   # use: typst 에 넘길 폰트 폴더(정적 한 쌍이 있으면 그 폴더만, 시스템 글꼴은 항상 무시)

# ---------------------------------------------------------------- typst 인라인 변환
_ESC = str.maketrans({c: "\\" + c for c in "\\#*_`$<>@[]~/\"'-+="})

def esc(t):
    """본문 글자를 typst 마크업 안에 안전하게 넣는다(특수문자 이스케이프)."""
    return str(t).translate(_ESC)

def _raw_str(t):
    return '"' + str(t).replace("\\", "\\\\").replace('"', '\\"') + '"'

_INL = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`|<sub>(.*?)</sub>|<sup>(.*?)</sup>|<br\s*/?>|<b>(.*?)</b>|<i>(.*?)</i>|<span[^>]*>(.*?)</span>|&nbsp;|&lt;|&gt;|&amp;", re.S)

def inline(text):
    """인라인 마크업(**굵게**, `코드`, <sub>, <sup>, <br>, <b>, <i>, <span>) → typst."""
    out, pos = [], 0
    for m in _INL.finditer(text):
        out.append(esc(text[pos:m.start()])); pos = m.end()
        g = m.groups(); tok = m.group(0)
        # 토큰 뒤의 ';' 는 typst 코드 표현식을 닫는 표시(출력되지 않음). 없으면 뒤따르는 '(' 가 함수 호출로 읽힌다(예: f<sub>P</sub>(N, α)).
        if g[0] is not None:   out.append(f"#strong[{inline(g[0])}];")
        elif g[1] is not None: out.append(f"#raw({_raw_str(g[1])});")
        elif g[2] is not None: out.append(f"#sub[{inline(g[2])}];")
        elif g[3] is not None: out.append(f"#super[{inline(g[3])}];")
        elif tok.startswith("<br"): out.append("#linebreak();")
        elif g[4] is not None: out.append(f"#strong[{inline(g[4])}];")
        elif g[5] is not None: out.append(f"#emph[{inline(g[5])}];")
        elif g[6] is not None: out.append(inline(g[6]))
        elif tok == "&nbsp;": out.append("~")
        elif tok == "&lt;": out.append("\\<")
        elif tok == "&gt;": out.append("\\>")
        elif tok == "&amp;": out.append("&")
    out.append(esc(text[pos:]))
    return "".join(out)

def _caption(text):
    """캡션 앞의 '표 1.' '그림 2.' 'Table 1.' 'Fig. 3.' 를 굵게."""
    m = re.match(r"^((?:표|그림|Table|Fig\.?)\s*\d+[.:])\s*(.*)$", text.strip(), re.S)
    return f"#strong[{inline(m.group(1))}] {inline(m.group(2))}" if m else inline(text)

# ---------------------------------------------------------------- typst 문서(블록 모음)
PREAMBLE = r'''#set page(paper: "a4", margin: (top: 20mm, bottom: 20mm, x: 18mm),
  footer: context align(center, text(size: 8pt, fill: rgb("#777"))[#counter(page).display()]))
#set text(font: ("__FAMILY__", "Noto Sans CJK KR", "DejaVu Sans"), size: 10pt, fill: rgb("#111"), lang: "ko")
#set par(leading: 0.6em, spacing: 6pt, justify: false)
#set list(marker: [•], indent: 6pt, body-indent: 5pt, spacing: 4pt)
#let h2(t) = block(sticky: true, above: 18pt, below: 6pt, width: 100%, fill: rgb("#f3f3f3"), stroke: (left: 4pt + rgb("#333")), inset: (left: 8pt, right: 4pt, y: 3pt), text(size: 12.5pt, weight: "bold", t))
#let h3(t) = block(sticky: true, above: 10pt, below: 3pt, text(size: 10.5pt, weight: "bold", t))
#let keybox(body) = block(width: 100%, stroke: 1.5pt + rgb("#333"), fill: rgb("#fbfbfb"), inset: (x: 12pt, y: 8pt), above: 8pt, below: 12pt, breakable: false, body)
#let two(a, b) = grid(columns: (1fr, 1fr), gutter: 14pt, a, b)
#let tcap(t) = align(center, text(size: 8.8pt, fill: rgb("#444"), t))
#let fcap(t) = align(center, text(size: 9pt, fill: rgb("#222"), t))
#let eq(t) = block(above: 5pt, below: 5pt, width: 100%, align(center, text(font: ("Libertinus Serif", "Times New Roman", "DejaVu Serif"), size: 10.5pt, t)))
#let fig(path, cap, w) = block(width: 100%, above: 8pt, below: 10pt, breakable: false, align(center, image(path, width: w)) + v(4pt) + fcap(cap))
#let tbl(body, cap) = block(width: 100%, above: 4pt, below: 8pt, breakable: false, align(center, text(size: 8.8pt, body)) + (if cap != none { v(2pt) + tcap(cap) } else { [] }))
'''

class Doc:
    """typst 조각과 블록 노드(Word 출력용)를 함께 쌓는다. 그림은 assets 에 모아 빌드 때 임시 폴더로 복사한다(경로 제약 없음)."""
    def __init__(self):
        self.parts, self.assets, self.plain = [], {}, []   # plain: 숫자·대시 점검용 원문 조각
        self.nodes = []                                     # (종류, …) 튜플. render_docx 가 읽는다
        self.col_mm = TEXT_W_MM                             # 이 Doc 이 담기는 칸 폭(2단 칸이면 TWO_COL_MM)
        self.fit = {"log": [], "left": [], "tmp": None}     # 그림 자동 맞춤 기록(중첩 Doc 과 공유)
    def add(self, typ, plain=None, node=None):
        self.parts.append(typ)
        if plain: self.plain.append(plain)
        if node: self.nodes.append(node)
        return self
    def asset(self, path):
        if re.match(r"^[a-z]+://", path):
            raise ValueError(f"원격 그림은 지원하지 않는다: {path}")
        path = os.path.abspath(path)
        if not os.path.exists(path):
            raise FileNotFoundError(f"그림 파일 없음: {path}")
        name = f"img{len(self.assets) + 1:02d}{os.path.splitext(path)[1].lower()}"
        self.assets[name] = path
        return name
    # ---- 블록
    def title(self, t, sub=None):
        self.add(f'#align(center, text(size: 17pt, weight: "bold", tracking: -0.3pt)[{inline(t)}])', "## " + t, ("title", t, sub))
        if sub: self.add(f'#v(2pt)\n#align(center, text(size: 9pt, fill: rgb("#555"))[{inline(sub)}])', "## " + sub)
        self.add("#v(12pt)")
        return self
    def h2(self, t): return self.add(f"#h2[{inline(t)}]", "## " + t, ("h2", t))
    def h3(self, t): return self.add(f"#h3[{inline(t)}]", "### " + t, ("h3", t))
    def p(self, t): return self.add(inline(t) + "\n", t, ("p", t))
    def ul(self, items): return self.add("\n".join("- " + inline(i) for i in items) + "\n", "\n".join(items), ("ul", list(items)))
    def eq(self, expr): return self.add(f"#eq[{inline(expr)}]", "$$" + expr + "$$", ("eq", expr))
    def pagebreak(self): return self.add("#pagebreak()", None, ("pagebreak",))
    def key(self, title, body_typ, plain="", nodes=()):
        return self.add(f"#keybox[#strong[{inline(title)}]\n{body_typ}]", plain, ("key", title, list(nodes)))
    def two(self, left_typ, right_typ, plain="", left_nodes=(), right_nodes=()):
        return self.add(f"#two[\n{left_typ}\n][\n{right_typ}\n]", plain, ("two", list(left_nodes), list(right_nodes)))
    def fig(self, path, caption, width=None, regen=None, fit=True, cwd=None):
        """그림 + 아래 캡션. fit=True(기본)면 rules.md W10 기준으로 자동 맞춤(흰 여백 자르기 → 폭 조정 →
        원천 재생성(regen) → 앞 문단 옆 배치) 후 남은 문제만 기록한다. regen 은 callable(width_mm, max_height_mm) → 새 PNG 경로
        또는 셸 명령 문자열({w_mm} {h_mm} {out} 치환)."""
        if not fit:
            name = self.asset(path); w = (width or "100%").strip()
            return self.add(f"#fig({_raw_str(name)}, [{_caption(caption)}], {w})", caption,
                            ("fig", self.assets[name], caption, _width_mm(w, self.col_mm) or self.col_mm))
        tag = f"{os.path.basename(path)} [{(caption or '').split('.')[0][:12]}]"
        path, w_mm, side_prev, log, left = autofit(self, os.path.abspath(path), width, regen, cwd)
        self.fit["log"] += [f"{tag} {m}" for m in log]; self.fit["left"] += [f"{tag} {m}" for m in left]
        name = self.asset(path)
        figt = f"fig({_raw_str(name)}, [{_caption(caption)}], {w_mm:.1f}mm)"
        fnode = ("fig", self.assets[name], caption, w_mm)
        if side_prev:   # 앞 문단과 나란히(그림은 오른쪽 칸)
            prev = self.parts.pop()
            pnode = self.nodes.pop() if self.nodes and self.nodes[-1][0] == "p" else ("p", "")
            return self.add(f"#grid(columns: (1fr, {w_mm:.1f}mm), gutter: 14pt, [\n{prev}\n], {figt})", caption,
                            ("side", pnode, fnode, w_mm))
        return self.add("#" + figt, caption, fnode)
    def table(self, header, rows, caption=None):
        n = len(header)
        cells = ["#table.hline(stroke: 2pt + rgb(\"#333\"))",
                 "table.header(" + ", ".join(f"[#text(weight: 600)[{inline(str(h))}]]" for h in header) + ")",
                 "table.hline(stroke: 1pt + rgb(\"#333\"))"]
        for k, r in enumerate(rows):
            r = list(r) + [""] * (n - len(r))
            cells.append(", ".join(f"[{inline(str(c))}]" for c in r[:n]))
            cells.append("table.hline(stroke: 0.5pt + rgb(\"#bbb\"))" if k < len(rows) - 1 else "table.hline(stroke: 2pt + rgb(\"#333\"))")
        body = (f"table(columns: {n}, stroke: none, align: center, inset: (x: 8pt, y: 3pt), "
                "fill: (x, y) => if y == 0 { rgb(\"#fafafa\") },\n  " + ",\n  ".join(cells).lstrip("#") + "\n)")
        cap = f"[{_caption(caption)}]" if caption else "none"
        plain = " | ".join(map(str, header)) + "\n" + "\n".join(" | ".join(map(str, r)) for r in rows) + ("\n" + caption if caption else "")
        return self.add(f"#tbl({body}, {cap})", plain, ("table", [str(h) for h in header], [[str(c) for c in r] for r in rows], caption))
    # ---- 출력
    def typst(self, family=FONT_FAMILY):
        return PREAMBLE.replace("__FAMILY__", family) + "\n" + "\n\n".join(self.parts) + "\n"
    def plain_text(self):
        return "\n".join(self.plain)

# ---------------------------------------------------------------- markdown → Doc
_TBL = re.compile(r"^\s*\|.*\|\s*$")
_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_IMG = re.compile(r"^!\[(.*?)\]\((.*?)\)(?:\{(.*)\})?\s*$")
_ATTR = re.compile(r'(\w+)=("[^"]*"|\S+)')
_CAP = re.compile(r"^(표|Table)\s*\d+[.:]")

def _split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]

def md_blocks(lines, base_dir, doc):
    """markdown 줄 목록을 doc 에 쌓는다. 중첩(상자·2단)은 임시 Doc 으로 변환해 typst 조각만 넘긴다."""
    i, n = 0, len(lines); para = []
    def flush():
        if para:
            doc.p(" ".join(s.strip() for s in para)); para.clear()
    def sub(body, col_mm=None):
        d = Doc(); d.assets = doc.assets; d.fit = doc.fit; d.col_mm = col_mm or doc.col_mm; md_blocks(body, base_dir, d)
        doc.plain.extend(d.plain); return "\n\n".join(d.parts), d.nodes
    while i < n:
        ln = lines[i]; s = ln.strip()
        if not s:
            flush(); i += 1; continue
        if s.startswith(":::"):
            flush(); kind, _, title = s[3:].strip().partition(" ")
            j = i + 1; body = []
            while j < n and lines[j].strip() != ":::":
                body.append(lines[j]); j += 1
            if kind == "key":
                typ, nodes = sub(body); doc.key(title or "핵심 결과", typ, nodes=nodes)
            elif kind == "two":
                k = next((t for t, b in enumerate(body) if b.strip() == "|||"), len(body))
                (lt, ln_), (rt, rn) = sub(body[:k], TWO_COL_MM), sub(body[k + 1:], TWO_COL_MM)
                doc.two(lt, rt, left_nodes=ln_, right_nodes=rn)
            else:
                typ, nodes = sub(body); doc.add(typ, None, ("group", nodes))
            i = j + 1; continue
        if s.startswith("### "): flush(); doc.h3(s[4:]); i += 1; continue
        if s.startswith("## "):  flush(); doc.h2(s[3:]); i += 1; continue
        if s.startswith("# "):   flush(); doc.title(s[2:]); i += 1; continue
        if s in ("\\pagebreak", "<!-- pagebreak -->"): flush(); doc.pagebreak(); i += 1; continue
        if s.startswith("$$") and s.endswith("$$") and len(s) > 4:
            flush(); doc.eq(s[2:-2].strip()); i += 1; continue
        m = _IMG.match(s)
        if m:
            flush(); cap, path = m.group(1), m.group(2)
            at = {k: v.strip('"') for k, v in _ATTR.findall(m.group(3) or "")}
            if not os.path.isabs(path):
                path = os.path.join(base_dir, path)
            doc.fig(path, cap, at.get("width"), regen=at.get("regen"), fit=at.get("fit", "on") != "off", cwd=base_dir)
            i += 1; continue
        if _TBL.match(s) and i + 1 < n and _SEP.match(lines[i + 1]):
            flush(); header = _split_row(s); rows = []; j = i + 2
            while j < n and _TBL.match(lines[j]):
                rows.append(_split_row(lines[j])); j += 1
            caption = None
            if j < n and _CAP.match(lines[j].strip()):
                caption = lines[j].strip(); j += 1
            doc.table(header, rows, caption); i = j; continue
        if s.startswith(("- ", "* ")):
            flush(); items = []
            while i < n and lines[i].strip().startswith(("- ", "* ")):
                item = lines[i].strip()[2:]; i += 1
                while i < n and lines[i].startswith("  ") and lines[i].strip() and not lines[i].strip().startswith(("- ", "* ")):
                    item += " " + lines[i].strip(); i += 1
                items.append(item)
            doc.ul(items); continue
        para.append(ln); i += 1
    flush()
    return doc

def parse_front_matter(text):
    meta = {}
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            for ln in text[3:end].strip().splitlines():
                k, _, v = ln.partition(":")
                if _: meta[k.strip()] = v.strip()
            text = text[end + 4:]
    return meta, text

def md_to_doc(md_text, base_dir="."):
    meta, body = parse_front_matter(md_text)
    doc = Doc()
    if meta.get("title"):
        doc.title(meta["title"], meta.get("subtitle"))
    md_blocks(body.splitlines(), base_dir, doc)
    return meta, doc

# ---------------------------------------------------------------- 점검
def check_text(plain):
    """산출 전 금지 문자 점검: em/en dash. 발견 문맥을 돌려준다."""
    bad = []
    for m in re.finditer("[\u2013\u2014]", plain):
        s = max(0, m.start() - 30); bad.append(plain[s:m.end() + 30].replace("\n", " "))
    return bad

def num(x, nd=3):
    """보고서 숫자 형식: 정수는 천 단위 쉼표, 실수는 소수 nd(기본 3)자리 반올림 + 천 단위 쉼표."""
    if isinstance(x, str):
        return x
    if float(x).is_integer() and not isinstance(x, float):
        return f"{int(x):,}"
    return f"{float(x):,.{nd}f}"

def figsize_mm(w_mm=174.0, h_mm=None, aspect=0.5):
    """matplotlib figsize(인치). 보고서에 넣을 실제 크기(mm)로 그리면 그림 안 글자 pt 가 지면 pt 와 같다(W10).
    예: plt.figure(figsize=figsize_mm(174, aspect=0.45)); regen 함수에서는 figsize_mm(w_mm, min(h_mm, w_mm * 0.5))."""
    h_mm = h_mm if h_mm is not None else w_mm * aspect
    return (w_mm / 25.4, h_mm / 25.4)

def check_numbers(plain, nd=3):
    """숫자 표기 점검: (1) 소수 1~(nd-1)자리로 끝난 값, (2) 쉼표 없는 4자리 이상 정수. 발견 문맥을 돌려준다.
    제외: `코드`, 장·절 제목 줄, 수식($$), <span class="nocheck">, 날짜 YYYY-MM-DD, 연도 4자리."""
    txt = re.sub(r"`[^`]*`|\$\$.*?\$\$|<span class=\"nocheck\">.*?</span>", " ", plain, flags=re.S)
    txt = "\n".join(l for l in txt.splitlines() if not l.startswith("#"))
    txt = re.sub(r"\b\d{4}-\d{2}-\d{2}\b", " ", txt)
    txt = re.sub(r"\d+\.\d+\s?절", " ", txt)                              # 절 번호(2.2절)
    txt = re.sub(r"<[^>]+>", " ", txt)
    bad = []
    for m in re.finditer(rf"(?<![\d.,])\d[\d,]*\.\d{{1,{nd - 1}}}(?![\d.])", txt):
        s = max(0, m.start() - 25); bad.append(f"소수 {nd}자리 아님: …{txt[s:m.end() + 15].strip()}…")
    for m in re.finditer(r"(?<![\d.,])\d{4,}(?![\d,])", txt):
        if re.fullmatch(r"(19|20)\d{2}", m.group()):
            continue
        s = max(0, m.start() - 25); bad.append(f"천 단위 쉼표 없음: …{txt[s:m.end() + 15].strip()}…")
    return bad

# ---------------------------------------------------------------- 그림 자동 맞춤(rules.md W10)
TEXT_W_MM = 174.0                      # A4 210 mm - 좌우 여백 18 mm × 2
TWO_COL_MM = (TEXT_W_MM - 4.94) / 2    # 2단 칸(간격 14 pt)
TEXT_H_MM = 257.0                      # A4 297 mm - 위아래 여백 20 mm × 2

def image_px(path):
    """PNG·JPEG 의 (가로, 세로) 픽셀. 다른 형식이나 읽기 실패는 None(외부 패키지 없이 헤더만 읽는다)."""
    try:
        with open(path, "rb") as f:
            head = f.read(26)
            if head[:8] == b"\x89PNG\r\n\x1a\n":
                return struct.unpack(">II", head[16:24])
            if head[:2] == b"\xff\xd8":
                f.seek(2)
                while True:
                    b = f.read(1)
                    while b and b != b"\xff": b = f.read(1)
                    while b == b"\xff": b = f.read(1)
                    if not b: return None
                    m = b[0]
                    if m in (0xD8, 0x01) or 0xD0 <= m <= 0xD7: continue
                    ln = struct.unpack(">H", f.read(2))[0]
                    if 0xC0 <= m <= 0xCF and m not in (0xC4, 0xC8, 0xCC):
                        h, w = struct.unpack(">xHH", f.read(5)); return (w, h)
                    f.seek(ln - 2, 1)
    except (OSError, struct.error):
        pass
    return None

def _width_mm(w, col_mm):
    m = re.fullmatch(r"\s*([\d.]+)\s*(%|mm|cm|in|pt)\s*", w or "100%")
    if not m: return None
    v, u = float(m.group(1)), m.group(2)
    return {"%": col_mm * v / 100, "mm": v, "cm": v * 10, "in": v * 25.4, "pt": v * 25.4 / 72}[u]

MIN_W = 0.60          # 좌우 빈 폭 40 % 이하 = 칸 폭의 60 % 이상
MAX_H = 0.75          # 그림 높이 ≤ 쪽 본문의 75 %
MIN_DPI = 150

def fig_problems(path, width, col_mm):
    """보고서 지면에 넣었을 때의 문제 목록(W10 판정 기준). 빈 목록이면 통과. width 는 '66%'·'80mm' 같은 문자열 또는 mm 수."""
    px = image_px(path); wmm = width if isinstance(width, (int, float)) else _width_mm(width, col_mm)
    if not px or not wmm: return []
    hmm = wmm * px[1] / px[0]; out = []
    side = 1 - wmm / col_mm
    if side > 1 - MIN_W + 1e-6:
        out.append((f"세로로 긴 그림을 좁게 넣음(좌우 빈 폭 {side * 100:.0f} %)" if px[1] > 1.2 * px[0]
                    else f"좌우 빈 폭 {side * 100:.0f} %"))
    if hmm > MAX_H * TEXT_H_MM + 1e-6:
        out.append(f"높이 {hmm:.0f} mm(쪽 본문의 {hmm / TEXT_H_MM * 100:.0f} %)")
    dpi = px[0] / (wmm / 25.4)
    if dpi < MIN_DPI - 1e-6:
        out.append(f"삽입 크기에서 {dpi:.0f} dpi({MIN_DPI} 미만)")
    return out

def natural_mm(path):
    """그림 파일에 기록된 dpi 로 계산한 실제 폭(mm). 기록이 없거나 Pillow 가 없으면 None."""
    try:
        from PIL import Image
        dpi = Image.open(path).info.get("dpi")
        return image_px(path)[0] / float(dpi[0]) * 25.4 if dpi and dpi[0] > 1 else None
    except Exception:
        return None

def _fit_tmp(doc):
    if not doc.fit["tmp"]:
        doc.fit["tmp"] = tempfile.mkdtemp(prefix="report_fit_")
    return doc.fit["tmp"]

def trim_white(path, tmpdir, thresh=6, pad_frac=0.01):
    """그림 둘레의 흰(또는 투명) 여백을 1 % 만 남기고 잘라 새 파일 경로를 돌려준다. 잘라낼 것이 없거나 Pillow 가 없으면 원래 경로."""
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return path
    try:
        im = Image.open(path); dpi = im.info.get("dpi")
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA"); bg = Image.new("RGBA", im.size, (255, 255, 255, 255)); bg.alpha_composite(im); im = bg
        rgb = im.convert("RGB")
        diff = ImageChops.difference(rgb, Image.new("RGB", rgb.size, (255, 255, 255))).convert("L").point(lambda v: 255 if v > thresh else 0)
        bbox = diff.getbbox()
        if not bbox: return path
        W, H = rgb.size; pad = max(2, round(pad_frac * max(W, H)))
        l, t, r, b = max(0, bbox[0] - pad), max(0, bbox[1] - pad), min(W, bbox[2] + pad), min(H, bbox[3] + pad)
        if (W - (r - l)) < 0.03 * W and (H - (b - t)) < 0.03 * H:
            return path
        out = os.path.join(tmpdir, "trim_" + os.path.splitext(os.path.basename(path))[0] + f"_{len(os.listdir(tmpdir))}.png")
        rgb.crop((l, t, r, b)).save(out, **({"dpi": dpi} if dpi else {}))
        return out
    except Exception:
        return path

def _run_regen(regen, w_mm, h_mm, tmpdir, cwd):
    """원천 재생성. callable 이면 (폭 mm, 최대 높이 mm) 로 부르고, 문자열이면 셸 명령으로 실행한다."""
    out = os.path.join(tmpdir, f"regen_{len(os.listdir(tmpdir))}.png")
    if callable(regen):
        res = regen(w_mm, h_mm)
        return os.path.abspath(res) if res else None
    cmd = regen.replace("{w_mm}", f"{w_mm:.1f}").replace("{h_mm}", f"{h_mm:.1f}").replace("{out}", out)   # 다른 중괄호는 그대로 둔다
    r = subprocess.run(cmd, shell=True, cwd=cwd or None, capture_output=True, text=True)
    if r.returncode != 0 or not os.path.exists(out):
        raise RuntimeError(f"regen 실패({r.returncode}): {cmd}\n{r.stderr[-400:]}")
    return out

def autofit(doc, path, width, regen=None, cwd=None):
    """rules.md W10 자동 맞춤. 돌려주는 값: (최종 경로, 폭 mm, 앞 문단 옆 배치 여부, 조치 기록, 남은 문제).
    순서: 흰 여백 자르기 → 문제가 있으면 원천 재생성(regen) → 폭 조정(늘릴 때는 기준선까지만, 글자 확대를 막는다)
    → 세로로 긴 그림은 앞 문단 옆 배치 → 그래도 남은 것은 '남은 문제'."""
    col = doc.col_mm; hmax = MAX_H * TEXT_H_MM; log = []
    tmp = _fit_tmp(doc)
    p2 = trim_white(path, tmp)
    if p2 != path:
        log.append("흰 여백 잘라냄"); path = p2
    px = image_px(path)
    w_req = _width_mm(width, col) if width else col
    if not px or not w_req:
        return path, w_req or col, False, log, []
    probs = fig_problems(path, w_req, col)
    if probs and regen:
        saved = (path, px, w_req)
        try:
            path = trim_white(_run_regen(regen, col, hmax, tmp, cwd), tmp); px = image_px(path)
            if not px or not px[0] or not px[1]:
                raise RuntimeError("재생성 결과를 읽지 못함")
            nat = natural_mm(path)            # 재생성 그림은 기록된 dpi 의 실제 크기 그대로(늘리면 글자가 커진다)
            w_req = min(col, hmax * px[0] / px[1], nat or col)
            log.append(f"원천 재생성({col:.0f} x ≤{hmax:.0f} mm 요청, {w_req:.0f} mm 로 넣음)")
            probs = fig_problems(path, w_req, col)
        except Exception as e:
            path, px, w_req = saved; probs = fig_problems(path, w_req, col)   # 원래 그림으로 되돌려 이어서 맞춘다
            log.append(f"재생성 실패: {e}".splitlines()[0])
    r = px[1] / px[0]
    if not probs and not (r > 1.2 and w_req * r > 0.5 * TEXT_H_MM):
        return path, w_req, False, log, []
    w = w_req
    w_hcap = hmax / r                                    # 높이 한도로 정해지는 최대 폭
    if w * r > hmax:
        w = w_hcap; log.append(f"높이 한도에 맞춰 폭 {w_req:.0f} → {w:.0f} mm")
    if w < MIN_W * col:
        w_new = min(MIN_W * col, w_hcap)
        if w_new > w:
            log.append(f"폭 {w:.0f} → {w_new:.0f} mm(좌우 빈 폭 40 % 기준선까지만)"); w = w_new
    side = False
    tall_big = r > 1.2 and w * r > 0.5 * TEXT_H_MM          # 세로 그림이 쪽 절반 넘게 차지(세로 사진 등)
    if (w < MIN_W * col - 1e-6 or tall_big) and col >= TEXT_W_MM - 1e-6 and doc.parts and _is_para(doc.parts[-1]):
        w = min(w_hcap, 0.5 * col); side = True
        log.append(f"세로로 긴 그림: 앞 문단 옆 배치(그림 칸 {w:.0f} mm)")
    dpi = px[0] / (w / 25.4)
    left = []
    if dpi < MIN_DPI:
        left.append(f"삽입 크기에서 {dpi:.0f} dpi({MIN_DPI} 미만): 원천에서 더 큰 해상도로 다시 내보내기")
    if not side:
        left += [m + ": 원천에서 보고서판으로 다시 그리기" for m in fig_problems(path, w, col) if "dpi" not in m]
    return path, w, side, log, left

def _is_para(part):
    """Doc.parts 조각이 일반 문단인가(블록 함수 호출이 아닌 글)."""
    return not part.lstrip().startswith("#") or part.lstrip().startswith(("#strong", "#emph", "#sub", "#super", "#raw", "#linebreak"))

def fig_report(doc):
    """자동 맞춤 기록과 남은 문제를 출력하고 남은 문제 수를 돌려준다."""
    if doc.fit["log"]:
        sys.stderr.write(f"[report_kit] 그림 자동 맞춤 {len(doc.fit['log'])}건(W10):\n  " + "\n  ".join(doc.fit["log"]) + "\n")
    if doc.fit["left"]:
        sys.stderr.write(f"[report_kit] 그림 배치 미해결 {len(doc.fit['left'])}건(자동으로 못 고침, W10 대로 원천 수정 필요):\n  "
                         + "\n  ".join(doc.fit["left"]) + "\n")
    return len(doc.fit["left"])

def report_checks(plain):
    bad = check_text(plain)
    if bad:
        sys.stderr.write("[report_kit] em/en dash 발견, 하이픈으로 바꾸시오:\n  " + "\n  ".join(bad) + "\n")
    warn = check_numbers(plain)
    if warn:
        sys.stderr.write(f"[report_kit] 숫자 표기 점검 {len(warn)}건(소수 셋째 자리·천 단위 쉼표 규칙; 예외는 <span class=\"nocheck\">):\n  " + "\n  ".join(warn) + "\n")
    return bad, warn

# ---------------------------------------------------------------- 빌드(typst)
def _compile(src, root, font_dirs, fmt=None, ppi=None):
    _use_venv_site()
    import typst
    main = os.path.join(root, "main.typ")
    open(main, "w", encoding="utf-8").write(src)
    kw = {"root": root, "font_paths": font_dirs, "ignore_system_fonts": True}   # 시스템 글꼴 무시: PC 가 달라도 같은 출력
    if fmt: kw["format"] = fmt
    if ppi: kw["ppi"] = ppi
    return typst.compile(main, **kw)

def _preview_from_src(src, root, font_dirs, out_dir, ppi=55):
    """쪽별 PNG(p01.png …) 와 접촉 시트(sheet1.png …, 4쪽씩). 모두 typst 로 그린다."""
    os.makedirs(out_dir, exist_ok=True)
    for f in glob.glob(os.path.join(out_dir, "*.png")):
        os.remove(f)
    pngs = _compile(src, root, font_dirs, fmt="png", ppi=ppi)
    pngs = pngs if isinstance(pngs, list) else [pngs]
    pages = []
    for k, b in enumerate(pngs):
        fp = os.path.join(out_dir, f"p{k + 1:02d}.png"); open(fp, "wb").write(b); pages.append(fp)
    sheets = []
    for k in range(0, len(pages), 4):
        chunk = pages[k:k + 4]
        sheet_src = ('#set page(width: auto, height: auto, margin: 0pt, fill: white)\n#grid(columns: %d, gutter: 0pt, %s)\n'
                     % (len(chunk), ", ".join(f'image({_raw_str(os.path.basename(p))}, width: 105mm)' for p in chunk)))
        b = _compile(sheet_src, out_dir, font_dirs, fmt="png", ppi=ppi * 2)
        b = b[0] if isinstance(b, list) else b
        sp = os.path.join(out_dir, f"sheet{k // 4 + 1}.png"); open(sp, "wb").write(b); sheets.append(sp)
    os.remove(os.path.join(out_dir, "main.typ"))
    return len(pages), sheets

def preview_md(md_path, out_dir=None, font=None, ppi=55):
    text = open(md_path, encoding="utf-8").read(); base = os.path.dirname(os.path.abspath(md_path))
    meta, doc = md_to_doc(text, base)
    _, _, _, dirs = font_status(font or meta.get("font"))
    out_dir = out_dir or os.path.splitext(md_path)[0] + "_preview"
    try:
        with tempfile.TemporaryDirectory() as td:
            for name, path in doc.assets.items():
                shutil.copy(path, os.path.join(td, name))
            return _preview_from_src(doc.typst(), td, dirs, out_dir, ppi)
    finally:
        if doc.fit["tmp"]:
            shutil.rmtree(doc.fit["tmp"], ignore_errors=True)

# ---------------------------------------------------------------- 빌드(Word .docx)
# 지면 값은 typst 전문(PREAMBLE)과 같다(rules.md F1·F3·F5). Word 는 반 pt 단위라 8.8 pt 는 9 pt 로 둔다.
# 글꼴은 맑은 고딕(받는 PC 에 늘 있다). Noto Sans KR 은 Windows 에 가변 폰트로만 깔린 경우가 많아 Word 에서 가늘고 굵게가 안 된다.
# 줄간격은 글꼴 줄 높이에 기대지 않게 pt 고정값으로 둔다(본문 10 pt × 1.6 = 16 pt).
DOCX_FAMILY = "Malgun Gothic"
DOCX_LINE_PT = 16

def inline_runs(text, fmt=None):
    """인라인 마크업 → [(글자, 서식 dict)]. 서식 키: bold, italic, sub, sup, code, br."""
    fmt = fmt or {}; out, pos = [], 0
    for m in _INL.finditer(text):
        if m.start() > pos: out.append((text[pos:m.start()], fmt))
        pos = m.end(); g = m.groups(); tok = m.group(0)
        if g[0] is not None:   out += inline_runs(g[0], {**fmt, "bold": True})
        elif g[1] is not None: out.append((g[1], {**fmt, "code": True}))
        elif g[2] is not None: out += inline_runs(g[2], {**fmt, "sub": True})
        elif g[3] is not None: out += inline_runs(g[3], {**fmt, "sup": True})
        elif tok.startswith("<br"): out.append(("", {**fmt, "br": True}))
        elif g[4] is not None: out += inline_runs(g[4], {**fmt, "bold": True})
        elif g[5] is not None: out += inline_runs(g[5], {**fmt, "italic": True})
        elif g[6] is not None: out += inline_runs(g[6], fmt)
        else: out.append(({"&nbsp;": " ", "&lt;": "<", "&gt;": ">", "&amp;": "&"}[tok], fmt))
    if pos < len(text): out.append((text[pos:], fmt))
    return out

def _text_mm(text, pt, bold=False):
    """글자 폭 어림(mm). 한글·한자 1 em, 그 밖 0.56 em, 굵게 5 % 더. <br> 로 나눈 줄 중 가장 긴 줄."""
    best = 0.0
    for line in re.split(r"<br\s*/?>", text):
        em = sum(1.0 if ord(ch) >= 0x1100 and not 0x2000 <= ord(ch) < 0x2E80 else 0.56
                 for s, f in inline_runs(line) for ch in s)
        best = max(best, em)
    return best * pt * 25.4 / 72 * (1.05 if bold else 1.0)

class _Docx:
    """python-docx 로 노드를 그린다. 컨테이너는 Document 또는 표 칸(_Cell)."""
    def __init__(self, family=DOCX_FAMILY):
        from docx import Document
        from docx.shared import Pt, Mm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        self.Pt, self.Mm, self.RGB, self.AL, self.BR, self.TA, self.OX, self.qn = Pt, Mm, RGBColor, WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TABLE_ALIGNMENT, OxmlElement, qn
        self.family = family
        self.d = d = Document()
        s = d.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin = s.bottom_margin = Mm(20); s.left_margin = s.right_margin = Mm(18)
        s.header_distance, s.footer_distance = Mm(8), Mm(8)
        st = d.styles["Normal"]; st.font.size = Pt(10); st.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        self._font(st.element.get_or_add_rPr(), family)
        pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(6); pf.line_spacing = Pt(DOCX_LINE_PT)
        lb = d.styles["List Bullet"].paragraph_format; lb.space_after = Pt(2)
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER   # 쪽번호: 아래 가운데 8 pt #777
        self._field(fp, "PAGE", size=8, color="777777")
    # ---- XML 도우미
    def _font(self, rPr, family, east=True):
        f = rPr.find(self.qn("w:rFonts"))
        if f is None:
            f = self.OX("w:rFonts"); rPr.insert(0, f)
        for k in ("w:ascii", "w:hAnsi", "w:cs") + (("w:eastAsia",) if east else ()):
            f.set(self.qn(k), family)
    def _el(self, tag, **attrs):
        e = self.OX(tag)
        for k, v in attrs.items(): e.set(self.qn("w:" + k), str(v))
        return e
    def _borders(self, parent_pr, tag, spec):
        """spec: {변: (굵기 1/8 pt, 색) 또는 None(없음)}"""
        b = parent_pr.find(self.qn(tag))
        if b is None:
            b = self.OX(tag); parent_pr.append(b)
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            if side not in spec: continue
            v = spec[side]
            b.append(self._el("w:" + side, val="nil") if v is None else self._el("w:" + side, val="single", sz=v[0], space=0, color=v[1]))
    def _shade(self, pr, fill):
        pr.append(self._el("w:shd", val="clear", color="auto", fill=fill))
    def _field(self, p, code, size=None, color=None):
        r = p.add_run(); self._fmt(r, size=size, color=color)
        r._r.append(self._el("w:fldChar", fldCharType="begin"))
        it = self.OX("w:instrText"); it.set(self.qn("xml:space"), "preserve"); it.text = f" {code} "; r._r.append(it)
        r._r.append(self._el("w:fldChar", fldCharType="separate"))
        t = self.OX("w:t"); t.text = "1"; r._r.append(t)
        r._r.append(self._el("w:fldChar", fldCharType="end"))
    def _fmt(self, r, size=None, color=None, bold=None):
        if size: r.font.size = self.Pt(size)
        if color: r.font.color.rgb = self.RGB.from_string(color)
        if bold is not None: r.bold = bold
    # ---- 글
    def runs(self, p, text, size=None, color=None, bold=None, serif=False):
        for s, f in inline_runs(text):
            if f.get("br"):
                p.add_run().add_break(); continue
            r = p.add_run(s); self._fmt(r, size, color, True if f.get("bold") else bold)
            if f.get("italic"): r.italic = True
            if f.get("sub"): r.font.subscript = True
            if f.get("sup"): r.font.superscript = True
            if f.get("code"): self._font(r._r.get_or_add_rPr(), "Consolas", east=False)
            elif serif: self._font(r._r.get_or_add_rPr(), "Times New Roman", east=False)
        return p
    def para(self, c, text="", align=None, size=None, color=None, bold=None, before=None, after=None, keep=False, style=None, line=None):
        p = c.add_paragraph(style=style) if style else c.add_paragraph()
        pf = p.paragraph_format
        if line: pf.line_spacing = line if isinstance(line, float) else self.Pt(line)   # float = 배수(그림 문단), int = pt 고정
        if align: p.alignment = align
        if before is not None: pf.space_before = self.Pt(before)
        if after is not None: pf.space_after = self.Pt(after)
        if keep: pf.keep_with_next = True
        return self.runs(p, text, size, color, bold) if text else p
    def caption(self, c, text, size, color, after):
        m = re.match(r"^((?:표|그림|Table|Fig\.?)\s*\d+[.:])\s*(.*)$", (text or "").strip(), re.S)
        p = self.para(c, align=self.AL.CENTER, before=2, after=after, line=13)
        if m:
            self.runs(p, m.group(1), size, color, True); self.runs(p, " " + m.group(2), size, color)
        else:
            self.runs(p, text, size, color)
        return p
    # ---- 표 공통
    def _table(self, c, rows, cols, width_mm=None, widths=None):
        t = c.add_table(rows=rows, cols=cols)
        tp = t._tbl.tblPr
        for e in tp.findall(self.qn("w:tblStyle")): tp.remove(e)
        if width_mm:
            t.autofit = False
            tw = tp.find(self.qn("w:tblW"))
            if tw is None:
                tw = self.OX("w:tblW"); tp.append(tw)
            tw.set(self.qn("w:type"), "dxa"); tw.set(self.qn("w:w"), str(round(width_mm / 25.4 * 1440)))
            for k, col in enumerate(t.columns):
                w = self.Mm(widths[k] if widths else width_mm / cols)
                col.width = w
                for cell in col.cells: cell.width = w
        return t
    def _cell_margins(self, t, x_pt, y_pt):
        m = self.OX("w:tblCellMar")
        for side, v in (("top", y_pt), ("left", x_pt), ("bottom", y_pt), ("right", x_pt)):
            m.append(self._el("w:" + side, w=round(v * 20), type="dxa"))
        t._tbl.tblPr.append(m)
    def _no_split(self, t):
        for row in t.rows:
            row._tr.get_or_add_trPr().append(self._el("w:cantSplit"))
    def _tidy_cell(self, cell):
        """칸 첫 빈 문단을 지우고, 칸이 표로 끝나면 빈 문단을 붙인다(Word 규칙)."""
        tc = cell._tc; kids = [e for e in tc if e.tag != self.qn("w:tcPr")]
        if len(kids) > 1 and kids[0].tag == self.qn("w:p") and not "".join(kids[0].itertext()).strip() \
                and kids[0].find(".//" + self.qn("w:drawing")) is None:
            tc.remove(kids[0])
        if tc[-1].tag == self.qn("w:tbl"):
            p = self.OX("w:p"); tc.append(p)
            ppr = self.OX("w:pPr"); p.append(ppr); ppr.append(self._el("w:spacing", after=0, line=20, lineRule="exact"))
    def spacer(self, c, pt):
        p = self.para(c, after=0); pf = p.paragraph_format; pf.line_spacing = self.Pt(pt)
        from docx.enum.text import WD_LINE_SPACING; pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        return p
    # ---- 블록
    def nodes(self, c, nodes, col_mm):
        for k, n in enumerate(nodes):
            getattr(self, "n_" + n[0])(c, n, col_mm)
            if n[0] == "p" and k + 1 < len(nodes) and nodes[k + 1][0] in ("table", "fig") and c.paragraphs:
                c.paragraphs[-1].paragraph_format.keep_with_next = True   # 표·그림을 설명 문단과 같은 쪽에(F3)
    def n_title(self, c, n, col):
        p = self.para(c, n[1], self.AL.CENTER, 17, "111111", True, after=2 if n[2] else 12, line=24)
        for r in p.runs: r.font._element.get_or_add_rPr().append(self._el("w:spacing", val=-6))   # 자간 -0.3 pt
        if n[2]: self.para(c, n[2], self.AL.CENTER, 9, "555555", after=12, line=13)
    def n_h2(self, c, n, col):
        p = self.para(c, n[1], size=12.5, bold=True, before=18, after=6, keep=True, line=21)
        pPr = p._p.get_or_add_pPr()
        self._borders(pPr, "w:pBdr", {"left": (32, "333333")})            # 왼쪽 4 pt 세로선
        pPr.find(self.qn("w:pBdr"))[0].set(self.qn("w:space"), "6")
        self._shade(pPr, "F3F3F3"); pPr.append(self._el("w:ind", left=140, right=80))
    def n_h3(self, c, n, col):
        self.para(c, n[1], size=10.5, bold=True, before=10, after=3, keep=True, line=16)
    def n_p(self, c, n, col):
        if n[1]: self.para(c, n[1])
    def n_ul(self, c, n, col):
        for it in n[1]: self.para(c, it, style="List Bullet")
    def n_eq(self, c, n, col):
        p = self.para(c, align=self.AL.CENTER, before=5, after=5, line=19); self.runs(p, n[1], size=10.5, serif=True)
    def n_pagebreak(self, c, n, col):
        self.para(c).add_run().add_break(self.BR.PAGE)
    def n_group(self, c, n, col):
        self.nodes(c, n[1], col)
    def n_key(self, c, n, col):
        t = self._table(c, 1, 1, col); t.alignment = self.TA.CENTER
        self._borders(t._tbl.tblPr, "w:tblBorders", {s: (12, "333333") for s in ("top", "left", "bottom", "right")})
        self._cell_margins(t, 12, 8); self._no_split(t)
        cell = t.cell(0, 0); self._shade(cell._tc.get_or_add_tcPr(), "FBFBFB")
        self.para(cell, f"**{n[1]}**", after=3, keep=True)
        self.nodes(cell, n[2], col - 8.5)
        for p in cell.paragraphs[:-1]: p.paragraph_format.keep_with_next = True
        self._tidy_cell(cell); self.spacer(c, 8)
    def _grid(self, c, widths, col):
        t = self._table(c, 1, len(widths), col, widths); t.alignment = self.TA.CENTER
        self._borders(t._tbl.tblPr, "w:tblBorders", {s: None for s in ("top", "left", "bottom", "right", "insideH", "insideV")})
        self._cell_margins(t, 0, 0)
        tl = self.OX("w:tblLayout"); tl.set(self.qn("w:type"), "fixed"); t._tbl.tblPr.append(tl)
        return t
    def n_two(self, c, n, col):
        gap = 4.94; w = (col - gap) / 2
        t = self._grid(c, [w, gap, w], col)
        for k, nodes in ((0, n[1]), (2, n[2])):
            self.nodes(t.cell(0, k), nodes, w); self._tidy_cell(t.cell(0, k))
        self.spacer(c, 4)
    def n_side(self, c, n, col):
        gap = 4.94; wf = n[3]; t = self._grid(c, [col - wf - gap, gap, wf], col)
        self.nodes(t.cell(0, 0), [n[1]], col - wf - gap); self.nodes(t.cell(0, 2), [n[2]], wf)
        for k in (0, 2): self._tidy_cell(t.cell(0, k))
        self.spacer(c, 4)
    def n_fig(self, c, n, col):
        p = self.para(c, align=self.AL.CENTER, before=8, after=0, keep=True, line=1.0)
        p.add_run().add_picture(n[1], width=self.Mm(min(n[3], col)))
        self.caption(c, n[2], 9, "222222", 10)
    def n_table(self, c, n, col):
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        header, rows, cap = n[1], n[2], n[3]; nc = len(header)
        body = [list(r) + [""] * (nc - len(r)) for r in rows]
        # 칸 폭은 내용에 맞춘다(typst columns: auto 와 같게). Word 자동 맞춤은 2단 칸 안에서 글자 단위로 무너지므로 직접 계산한다
        pad = 2 * 8 * 25.4 / 72 + 1.0
        widths = [max(_text_mm(str(v), 9, i == 0) for i, v in enumerate([header[j]] + [r[j] for r in body])) + pad for j in range(nc)]
        if sum(widths) > col:
            widths = [w * col / sum(widths) for w in widths]
        t = self._table(c, 1 + len(rows), nc, sum(widths), widths); t.alignment = self.TA.CENTER
        nxt = t._tbl.getnext()                                             # python-docx 가 칸 안 표 뒤에 붙이는 빈 문단
        if nxt is not None and nxt.tag == self.qn("w:p") and not "".join(nxt.itertext()).strip():
            nxt.getparent().remove(nxt)
        tp = t._tbl.tblPr
        tl = self.OX("w:tblLayout"); tl.set(self.qn("w:type"), "fixed"); tp.append(tl)
        self._borders(tp, "w:tblBorders", {"top": (16, "333333"), "bottom": (16, "333333"), "insideH": (4, "BBBBBB"),
                                            "left": None, "right": None, "insideV": None})
        self._cell_margins(t, 8, 3); self._no_split(t)
        t.rows[0]._tr.get_or_add_trPr().append(self._el("w:tblHeader"))
        for i, vals in enumerate([header] + body):
            for j in range(nc):
                cell = t.cell(i, j); p = cell.paragraphs[0]; p.alignment = self.AL.CENTER
                pf = p.paragraph_format; pf.space_after = self.Pt(0); pf.line_spacing = self.Pt(12.5); pf.keep_with_next = True
                self.runs(p, str(vals[j]), size=9, bold=True if i == 0 else None)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr(); self._shade(tcPr, "FAFAFA")
                    self._borders(tcPr, "w:tcBorders", {"top": (16, "333333"), "bottom": (8, "333333")})
        if cap:
            self.caption(c, cap, 9, "444444", 8)
        else:
            self.spacer(c, 6)
    def save(self, out):
        self.d.save(out)

def render_docx(doc, out_docx, family=DOCX_FAMILY, cleanup=True):
    """Doc → Word(.docx). 그림은 docx 안에 내장한다. 글꼴은 이름만 지정한다(기본 맑은 고딕, Windows 기본 글꼴)."""
    w = _Docx(family)
    w.nodes(w.d, doc.nodes, TEXT_W_MM)
    os.makedirs(os.path.dirname(os.path.abspath(out_docx)) or ".", exist_ok=True)
    w.save(out_docx)
    if cleanup and doc.fit["tmp"]:
        shutil.rmtree(doc.fit["tmp"], ignore_errors=True); doc.fit["tmp"] = None
    return out_docx

# ---- Word 미리보기: docx → PDF(Word 또는 LibreOffice) → 쪽 PNG
_PS = "/mnt/c/Windows/System32/WindowsPowerShell/v1.0/powershell.exe"

def _word_pdf(docx, pdf):
    """설치된 Microsoft Word 로 PDF 를 내보낸다(Windows, 또는 WSL 에서 Windows 의 Word). 실패하면 이유 문자열을 돌려준다."""
    ps = "powershell" if os.name == "nt" else (_PS if os.path.exists(_PS) else None)
    if not ps:
        return "Word 없음"
    script = ("$ErrorActionPreference='Stop'; $w=New-Object -ComObject Word.Application; $w.Visible=$false; $w.DisplayAlerts=0; "
              "try {{ $d=$w.Documents.Open('{i}', $false, $true); $d.Fields.Update() | Out-Null; $d.ExportAsFixedFormat('{o}', 17); $d.Close(0) }} finally {{ $w.Quit() }}")
    try:
        if os.name == "nt":
            r = subprocess.run([ps, "-NoProfile", "-Command", script.format(i=os.path.abspath(docx), o=os.path.abspath(pdf))], capture_output=True, text=True, timeout=180)
            return None if r.returncode == 0 and os.path.exists(pdf) else (r.stderr or r.stdout).strip()[-200:]
        wtmp = subprocess.run([ps, "-NoProfile", "-Command", "$env:TEMP"], capture_output=True, text=True, timeout=60).stdout.strip()
        ltmp = subprocess.run(["wslpath", "-u", wtmp], capture_output=True, text=True).stdout.strip()
        if not ltmp or not os.path.isdir(ltmp):
            return "Windows 임시 폴더를 찾지 못함"
        work = tempfile.mkdtemp(prefix="report_kit_", dir=ltmp); wwork = wtmp + "\\" + os.path.basename(work)
        try:
            shutil.copy(docx, os.path.join(work, "in.docx"))
            r = subprocess.run([ps, "-NoProfile", "-Command", script.format(i=wwork + "\\in.docx", o=wwork + "\\out.pdf")], capture_output=True, text=True, timeout=180)
            if not os.path.exists(os.path.join(work, "out.pdf")):
                return (r.stderr or r.stdout).strip()[-200:] or "Word 변환 실패"
            shutil.copy(os.path.join(work, "out.pdf"), pdf); return None
        finally:
            shutil.rmtree(work, ignore_errors=True)
    except (OSError, subprocess.SubprocessError) as e:
        return f"{type(e).__name__}: {e}"

def _soffice_pdf(docx, pdf):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return "LibreOffice 없음"
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", td, docx], capture_output=True, text=True, timeout=180)
        out = os.path.join(td, os.path.splitext(os.path.basename(docx))[0] + ".pdf")
        if not os.path.exists(out):
            return (r.stderr or r.stdout).strip()[-200:] or "LibreOffice 변환 실패"
        shutil.copy(out, pdf)
    return None

def docx_to_pdf(docx, pdf):
    """Word → LibreOffice 순으로 시도한다. 돌려주는 값: (사용한 엔진 이름 또는 None, 실패 이유 목록)."""
    why = []
    for name, fn in (("Word", _word_pdf), ("LibreOffice", _soffice_pdf)):
        err = fn(docx, pdf)
        if err is None:
            return name, why
        why.append(f"{name}: {err}")
    return None, why

def _pdf_pngs(pdf, out_dir, ppi):
    """PDF 쪽 → p01.png …(PyMuPDF, 없으면 typst 의 PDF 이미지)."""
    pages = []
    try:
        import fitz
        for k, pg in enumerate(fitz.open(pdf)):
            fp = os.path.join(out_dir, f"p{k + 1:02d}.png"); pg.get_pixmap(dpi=ppi).save(fp); pages.append(fp)
        return pages
    except ImportError:
        pass
    _use_venv_site(); import typst
    shutil.copy(pdf, os.path.join(out_dir, "src.pdf"))
    n = 0
    while True:
        n += 1; src = f'#set page(width: auto, height: auto, margin: 0pt)\n#image("src.pdf", page: {n}, width: 210mm)\n'
        main = os.path.join(out_dir, "page.typ"); open(main, "w", encoding="utf-8").write(src)
        try:
            b = typst.compile(main, root=out_dir, format="png", ppi=ppi)
        except Exception:
            break
        fp = os.path.join(out_dir, f"p{n:02d}.png"); open(fp, "wb").write(b[0] if isinstance(b, list) else b); pages.append(fp)
    for f in ("page.typ", "src.pdf"):
        if os.path.exists(os.path.join(out_dir, f)): os.remove(os.path.join(out_dir, f))
    return pages

def _sheets(pages, out_dir, ppi):
    """쪽 PNG 를 4쪽씩 이은 접촉 시트(Pillow). 없으면 빈 목록."""
    try:
        from PIL import Image
    except ImportError:
        return []
    out = []
    for k in range(0, len(pages), 4):
        ims = [Image.open(p).convert("RGB") for p in pages[k:k + 4]]
        h = max(i.height for i in ims); sheet = Image.new("RGB", (sum(i.width for i in ims), h), "white"); x = 0
        for i in ims:
            sheet.paste(i, (x, 0)); x += i.width
        sp = os.path.join(out_dir, f"sheet{k // 4 + 1}.png"); sheet.save(sp); out.append(sp)
    return out

def preview_docx(docx, out_dir, ppi=55):
    """Word 로 실제 쪽을 그려 미리보기를 만든다. 돌려주는 값: (쪽 수, 시트 목록, 엔진 이름) 또는 엔진이 없으면 (None, [], 이유)."""
    os.makedirs(out_dir, exist_ok=True)
    for f in glob.glob(os.path.join(out_dir, "*.png")) + glob.glob(os.path.join(out_dir, "*.pdf")):
        os.remove(f)
    pdf = os.path.join(out_dir, "word.pdf")
    eng, why = docx_to_pdf(docx, pdf)
    if not eng:
        return None, [], "; ".join(why)
    pages = _pdf_pngs(pdf, out_dir, ppi * 2); os.remove(pdf)
    return len(pages), _sheets(pages, out_dir, ppi), eng

def build_docx(doc, out, font=None, preview_dir=None, ppi=55):
    """Word(.docx)로 빌드한다. 돌려주는 값: (경로, 쪽 수 또는 None, 시트 목록). font 는 근사 미리보기용 폰트."""
    out = os.path.splitext(out)[0] + ".docx"
    render_docx(doc, out, cleanup=False)
    try:
        return (out,) + _docx_preview(doc, out, font, preview_dir, ppi)
    finally:
        if doc.fit["tmp"]:
            shutil.rmtree(doc.fit["tmp"], ignore_errors=True); doc.fit["tmp"] = None

def _docx_preview(doc, out, font, preview_dir, ppi):
    if preview_dir is None:
        return None, []
    pdir = preview_dir or os.path.splitext(out)[0] + "_preview"
    n, sheets, eng = preview_docx(out, pdir, ppi)
    if n is None:   # Word·LibreOffice 가 없으면 같은 원고의 typst 조판으로 근사 미리보기
        sys.stderr.write(f"[report_kit] Word 미리보기 불가({eng}). typst 근사 미리보기로 대신한다(쪽 나눔이 Word 와 다를 수 있다).\n")
        _, _, _, dirs = font_status(font)
        with tempfile.TemporaryDirectory() as td:
            for name, path in doc.assets.items():
                shutil.copy(path, os.path.join(td, name))
            n, sheets = _preview_from_src(doc.typst(), td, dirs, pdir, ppi)
    else:
        print(f"preview engine: {eng}")
    return n, sheets

# ---------------------------------------------------------------- Python API
class Report:
    """블록을 순서대로 쌓고 build() 로 Word(.docx). 문자열 인자는 markdown 인라인(**굵게**, `코드`, <sub>, <br>) 을 받는다."""
    def __init__(self, title, subtitle=""):
        self.doc = Doc(); self.doc.title(title, subtitle or None)
    def h2(self, t): self.doc.h2(t); return self
    def h3(self, t): self.doc.h3(t); return self
    def p(self, t): self.doc.p(t); return self
    def ul(self, items): self.doc.ul(items); return self
    def key(self, title, items):
        d = Doc(); d.ul(items); self.doc.key(title, "\n\n".join(d.parts), "\n".join(items), d.nodes); return self
    def table(self, header, rows, caption=None, nd=3):
        rows = [[num(c, nd) for c in r] for r in rows]; self.doc.table(header, rows, caption); return self
    def fig(self, path, caption, width=None, regen=None, fit=True):
        self.doc.fig(path, caption, width, regen=regen, fit=fit); return self
    def two(self, left_md, right_md):
        l, r = Doc(), Doc(); l.assets = r.assets = self.doc.assets; l.fit = r.fit = self.doc.fit; l.col_mm = r.col_mm = TWO_COL_MM
        md_blocks(left_md.splitlines(), ".", l); md_blocks(right_md.splitlines(), ".", r)
        self.doc.two("\n\n".join(l.parts), "\n\n".join(r.parts), l.plain_text() + "\n" + r.plain_text(), l.nodes, r.nodes); return self
    def eq(self, expr): self.doc.eq(expr); return self
    def pagebreak(self): self.doc.pagebreak(); return self
    def build(self, out, preview_dir=None, strict=None, font=None):
        """Word(.docx) 를 쓴다. strict(기본: 환경변수 REPORT_STRICT=1 이면 참)면
        그림 배치 미해결이 있을 때 파일을 쓴 뒤 종료 코드 3 으로 끝낸다."""
        report_checks(self.doc.plain_text()); left = fig_report(self.doc)
        out, n, sheets = build_docx(self.doc, out, font, preview_dir)
        if n is not None:
            print(f"pages: {n}; preview: {', '.join(sheets)}")
        print("saved:", out)
        if left and (strict if strict is not None else os.environ.get("REPORT_STRICT") == "1"):
            sys.exit(3)
        return out

# ---------------------------------------------------------------- 환경 점검·준비
def _has(mod):
    import importlib.util
    return importlib.util.find_spec(mod) is not None

def _placeholder_png(path, w=640, h=360, gray=200):
    """matplotlib 없이 회색 PNG 를 만든다(자기 검증용)."""
    raw = b"".join(b"\x00" + bytes([gray]) * (w * 3) for _ in range(h))
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
    open(path, "wb").write(png)

def _venv_python():
    for p in (os.path.join(VENV_DIR, "bin", "python"), os.path.join(VENV_DIR, "Scripts", "python.exe")):
        if os.path.exists(p):
            return p
    return None

def _reexec_if_needed():
    """typst 가 이 파이썬에 없고 예비 venv 에 있으면 그 파이썬으로 다시 실행한다. 사용자는 늘 python3 로 부르면 된다."""
    if (_has("typst") and _has("docx")) or os.environ.get("REPORT_KIT_NO_REEXEC"):
        return
    vp = _venv_python()
    if vp and os.path.abspath(vp) != os.path.abspath(sys.executable):
        os.environ["REPORT_KIT_NO_REEXEC"] = "1"
        os.execv(vp, [vp] + sys.argv)

def _use_venv_site():
    """Python API 로 import 됐을 때: typst 가 없고 예비 venv 가 있으면 그 site-packages 를 경로에 더한다."""
    if _has("typst") and _has("docx"):
        return True
    for sp in glob.glob(os.path.join(VENV_DIR, "lib", "python*", "site-packages")) + [os.path.join(VENV_DIR, "Lib", "site-packages")]:
        if os.path.isdir(sp):
            import site; site.addsitedir(sp)
    return _has("typst")

def _install_typst():
    """1) pip install typst  2) PEP 668(externally-managed) 이면 --user --break-system-packages(사용자 폴더 ~/.local 에만 설치)
    3) 그래도 안 되면 사용자 데이터 폴더에 venv 를 만들어 설치(이후 실행은 자동으로 그 파이썬을 쓴다). 관리자 권한은 어느 단계에도 없다."""
    def run(cmd):
        r = subprocess.run(cmd, capture_output=True, text=True); msg = (r.stderr or r.stdout).strip().splitlines()
        return r.returncode == 0, (msg[-1][:140] if msg else ""), (r.stderr + r.stdout)
    pkgs = [p for p, m in (("python-docx", "docx"), ("typst", "typst")) if not _has(m)] or ["python-docx", "typst"]
    base = [sys.executable, "-m", "pip", "install", "--quiet"] + pkgs
    ok, last, full = run(base)
    if ok:
        print(f"pip: {' '.join(pkgs)} 설치 완료(현재 파이썬)"); return True
    if "externally-managed" in full:
        ok, last, _ = run(base + ["--user", "--break-system-packages"])
        if ok:
            print("pip: 시스템 파이썬이 외부 관리형(PEP 668)이라 사용자 폴더(~/.local)에 설치했다. 시스템 파일은 건드리지 않는다"); return True
    print("pip: 현재 파이썬에는 설치하지 못했다 →", last)
    print(f"venv: {VENV_DIR} 를 만들어 그곳에 설치한다")
    ok, last, _ = run([sys.executable, "-m", "venv", VENV_DIR])
    if not ok:
        print("venv: 실패 →", last, "\n      python3-venv 가 없는 Debian/Ubuntu 라면 `sudo apt install python3-venv` 후 다시 setup"); return False
    ok, last, _ = run([_venv_python(), "-m", "pip", "install", "--quiet", "python-docx", "typst"])
    if not ok:
        print("venv: pip 실패 →", last); return False
    print("venv: 설치 완료. 이후 report_kit.py 는 자동으로 이 venv 파이썬으로 실행된다"); return True

def selftest():
    """예제 보고서를 임시 폴더에 Word 로 실제로 빌드해 표·그림이 들어갔는지 확인한다. 문제 수를 돌려준다."""
    exd = os.path.join(SKILL_DIR, "examples"); ex = os.path.join(exd, "example_report.md"); mk = os.path.join(exd, "make_example_figs.py")
    if not os.path.exists(ex):
        print("  [--] selftest    examples/example_report.md 없음"); return 1
    with tempfile.TemporaryDirectory() as td:
        try:
            if _has("matplotlib") and os.path.exists(mk):
                subprocess.run([sys.executable, mk, td], check=True, capture_output=True)
            else:
                for f in ("fig1_curves.png", "fig2_panels.png"):
                    _placeholder_png(os.path.join(td, f))
            text = open(ex, encoding="utf-8").read()
            meta, doc = md_to_doc(text, td)
            out = os.path.join(td, "selftest.docx"); render_docx(doc, out)
            import zipfile
            with zipfile.ZipFile(out) as z:
                xml = z.read("word/document.xml").decode("utf-8"); media = [f for f in z.namelist() if f.startswith("word/media/")]
            n_tbl = xml.count("<w:tbl>")
            ok_docx = n_tbl >= 1 and len(media) >= 1 and "핵심 결과" in xml
            print(f"  [{'ok' if ok_docx else '--'}] selftest    예제 Word 빌드: 표(상자·2단 포함) {n_tbl}개, 그림 {len(media)}개")
            return 0 if ok_docx else 1
        except Exception as e:
            print(f"  [--] selftest    빌드 실패: {type(e).__name__}: {str(e).splitlines()[0][:150]}"); return 1

def _selftest_preview():
    """Word(또는 LibreOffice)로 작은 docx 를 PDF 로 바꿔 미리보기 엔진을 확인한다. 실패해도 준비 상태에는 넣지 않는다(근사 미리보기가 있다)."""
    with tempfile.TemporaryDirectory() as td:
        d = Doc(); d.title("미리보기 시험", "report_kit · 2026-01-01"); d.p("한글 test 1,234.567")
        out = os.path.join(td, "t.docx"); render_docx(d, out)
        eng, why = docx_to_pdf(out, os.path.join(td, "t.pdf"))
        print(f"  [{'ok' if eng else '  '}] preview     " + (f"{eng} 로 docx 쪽 그리기 통과" if eng else "변환 실패: " + "; ".join(why)[:150]))

def doctor(build=False):
    """python-docx·typst·폰트·Word 상태를 출력하고 빠진 항목 수를 돌려준다. build=True 면 예제를 실제로 빌드해 본다."""
    n_bad = 0
    print(f"skill dir : {SKILL_DIR}\npython    : {sys.executable}")
    _use_venv_site(); ok_docx, ok_typst = _has("docx"), _has("typst")
    print(f"  [{'ok' if ok_docx else '--'}] python-docx (Word 산출물)" + ("" if ok_docx else "   설치: python3 -m pip install python-docx"))
    n_bad += 0 if ok_docx else 1
    word = (os.name == "nt" or os.path.exists(_PS)) and "Word(설치 여부는 --build 에서 확인)"
    eng = word or ("LibreOffice" if shutil.which("soffice") or shutil.which("libreoffice") else None)
    print(f"  [{'ok' if eng else '  '}] preview     Word 쪽 미리보기 엔진: {eng or '없음(typst 근사 미리보기로 대신)'}")
    # typst·폰트는 Word·LibreOffice 가 없을 때의 근사 미리보기에만 쓴다. 그때만 필수로 센다
    need = not eng; mark = "--" if need else "  "
    print(f"  [{'ok' if ok_typst else mark}] typst       (근사 미리보기, Word 가 없을 때만 필요)" + ("" if ok_typst else "   설치: python3 -m pip install typst"))
    n_bad += 1 if need and not ok_typst else 0
    reg, bold, var, dirs = font_status()
    if reg and bold:
        print(f"  [ok] font        {os.path.dirname(reg)}  (Regular·Bold, 근사 미리보기용)")
    else:
        print(f"  [{mark}] font        근사 미리보기용 정적 Noto Sans KR 없음"); n_bad += 1 if need else 0
    print(f"  [{'ok' if _has('matplotlib') else '  '}] matplotlib  (선택: 예제·그림 스크립트에만 필요)")
    if ok_typst and reg and bold:
        try:
            with tempfile.TemporaryDirectory() as td:
                _compile('#set text(font: "%s")\n한글 test 1,234.567' % FONT_FAMILY, td, dirs)
            print("  [ok] render      근사 미리보기 시험 렌더 통과")
        except Exception as e:
            print(f"  [{mark}] render      근사 미리보기 실패: {type(e).__name__}: {str(e).splitlines()[0][:120]}"); n_bad += 1 if need else 0
    if build and n_bad == 0:
        n_bad += selftest()
        if eng:
            _selftest_preview()
    if n_bad:
        print(f"준비: python3 {os.path.abspath(__file__)} setup")
    print("상태: " + ("준비됨" if n_bad == 0 else f"빠진 항목 {n_bad}개"))
    return n_bad

def setup(do_pip=True, do_font=True):
    """python-docx·typst(근사 미리보기용)가 없으면 pip 로 설치하고, 정적 한글 폰트(Regular·Bold)가 어디에도 없으면 사용자 데이터 폴더(DATA_DIR/fonts)에 내려받는다. 관리자 권한 불필요."""
    if do_pip and not (_has("typst") and _has("docx")):
        _install_typst()
        if not (_has("typst") and _has("docx")) and _venv_python():   # venv 에 설치됐으면 그 파이썬으로 이어서 점검한다
            os.environ["REPORT_KIT_NO_REEXEC"] = "1"
            os.execv(_venv_python(), [_venv_python()] + sys.argv + (["--no-pip"] if "--no-pip" not in sys.argv else []))
    if do_font:
        reg, bold, var, _ = font_status()
        if not (reg and bold):
            os.makedirs(USER_FONT_DIR, exist_ok=True)
            import urllib.request
            for name, url in FONT_FILES.items():
                dst = os.path.join(USER_FONT_DIR, name)
                if os.path.exists(dst):
                    continue
                print("font: 내려받기", url); urllib.request.urlretrieve(url, dst); print("font: 저장", dst)
    return doctor()

# ---------------------------------------------------------------- CLI
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="markdown → Word(.docx)")
    b.add_argument("md"); b.add_argument("-o", "--out")
    b.add_argument("--preview", nargs="?", const="", help="쪽 미리보기 PNG 폴더(생략 시 <out>_preview)")
    b.add_argument("--lenient", action="store_true", help="그림 배치 미해결이 남아도 종료 코드 0")
    v = sub.add_parser("preview", help=".docx → 쪽 PNG(Word 로 그림) 또는 markdown → 쪽 PNG(typst 근사, 빌드 없이)")
    v.add_argument("md"); v.add_argument("--dir"); v.add_argument("--font"); v.add_argument("--ppi", type=int, default=55)
    dc = sub.add_parser("doctor", help="python-docx·typst·폰트·Word 점검"); dc.add_argument("--build", action="store_true", help="예제 보고서를 임시 폴더에 빌드해 검증")
    st = sub.add_parser("setup", help="python-docx·typst(근사 미리보기용) 설치 + 한글 폰트 준비")
    st.add_argument("--no-pip", action="store_true"); st.add_argument("--no-font", action="store_true")
    a = ap.parse_args()
    if a.cmd != "setup":
        _reexec_if_needed()
    if a.cmd == "doctor":
        sys.exit(1 if doctor(build=a.build) else 0)
    if a.cmd == "setup":
        sys.exit(1 if setup(not a.no_pip, not a.no_font) else 0)
    if a.cmd == "preview":
        if a.md.lower().endswith(".docx"):
            n, sheets, eng = preview_docx(a.md, a.dir or os.path.splitext(a.md)[0] + "_preview", a.ppi)
            if n is None:
                sys.exit(f"[report_kit] Word 미리보기 불가: {eng}")
            print(f"preview engine: {eng}")
        else:
            n, sheets = preview_md(a.md, a.dir, a.font, a.ppi)
        print(f"pages: {n}; preview: {', '.join(sheets)}"); return
    text = open(a.md, encoding="utf-8").read()
    base = os.path.dirname(os.path.abspath(a.md))
    meta, doc = md_to_doc(text, base)
    if a.out:                                                    # -o 는 현재 디렉터리 기준
        out = os.path.abspath(a.out)
    else:                                                        # 머리말 output 은 원고 파일 폴더 기준
        out = meta.get("output") or os.path.splitext(os.path.basename(a.md))[0] + ".docx"
        if not os.path.isabs(out):
            out = os.path.join(base, out)
    report_checks(doc.plain_text()); left = fig_report(doc)
    out, n, sheets = build_docx(doc, out, meta.get("font"), a.preview if a.preview is not None else None)
    print("saved:", out)
    if n is not None:
        print(f"pages: {n}; preview: {', '.join(sheets)}")
    if left and not a.lenient:
        sys.exit(3)   # 파일은 썼지만 그림 배치 미해결이 남음(W10 대로 원천 수정 후 다시 빌드)

if __name__ == "__main__":
    main()
